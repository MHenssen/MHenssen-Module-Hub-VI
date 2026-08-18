"""Excel workbook I/O: the styled data workbook is the source of truth."""
import json
import datetime as dt

from .config import DEFAULT_DATA_XLSX, app_path
from .fields import _portfolio_product, _excel_exact_closed


def data_excel_path():
    return app_path(DEFAULT_DATA_XLSX)


def write_portfolio_excel(data, path=None):
    """Write the Jira portfolio cache as a styled, human-readable workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    path = path or data_excel_path()
    wb = Workbook()
    wb.remove(wb.active)
    meta = dict((data or {}).get("meta") or {})
    epics = list((data or {}).get("epics") or [])

    # Workbook palette, matching the dashboard theme.
    C = {"ink": "23262E", "purple": "5B50E6", "orange": "F48025", "green": "16A87A",
         "amber": "C8821F", "light": "F6F5F1", "line": "E4E2DA", "white": "FFFFFF",
         "red": "D64545", "soft_green": "DEF3EA", "soft_amber": "F8EDDC"}
    thin = Side(style="thin", color=C["line"])

    def setup_sheet(name, title, subtitle, headers, widths, tab_color=None):
        ws = wb.create_sheet(name)
        if tab_color:
            ws.sheet_properties.tabColor = tab_color
        ws.sheet_view.showGridLines = False
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))
        ws.cell(1, 1, title)
        ws.cell(1, 1).font = Font(name="Segoe UI", size=16, bold=True, color=C["white"])
        ws.cell(1, 1).fill = PatternFill("solid", fgColor=C["ink"])
        ws.cell(1, 1).alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 28
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(1, len(headers)))
        ws.cell(2, 1, subtitle)
        ws.cell(2, 1).font = Font(name="Segoe UI", size=9, color="6A6F7B")
        ws.row_dimensions[2].height = 24
        for c, h in enumerate(headers, 1):
            cell = ws.cell(4, c, h)
            cell.font = Font(name="Segoe UI", size=9, bold=True, color=C["white"])
            cell.fill = PatternFill("solid", fgColor=C["purple"])
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[4].height = 30
        ws.freeze_panes = "A5"
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return ws

    # ---------------------------------------------------------------- Read Me
    ws = wb.create_sheet("Read Me")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = C["orange"]
    ws.merge_cells("A1:D1")
    ws["A1"] = "PI Portfolio Explorer"
    ws["A1"].font = Font(name="Segoe UI", size=18, bold=True, color=C["white"])
    ws["A1"].fill = PatternFill("solid", fgColor=C["ink"])
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    rows = [
        ("Purpose", "All-team Parent Link and Epic progress data synchronized from Jira."),
        ("Target PI", meta.get("piName", "")),
        ("Synchronized", meta.get("syncedAt", "")),
        ("Jira", meta.get("jiraUrl", "")),
        ("Kanban board", meta.get("kanbanBoardId", "")),
        ("Teams", ", ".join(meta.get("teams") or [])),
        ("Dashboard", "The local dashboard reads this workbook directly."),
        ("Sheets", "Epics = one row per Epic; Tickets = child detail; Commitment = readable historical baseline; Parent Links and Teams = aggregated overviews."),
    ]
    for r, (label, value) in enumerate(rows, 3):
        ws.cell(r, 1, label).font = Font(name="Segoe UI", bold=True, color=C["ink"])
        ws.cell(r, 1).fill = PatternFill("solid", fgColor=C["light"])
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.cell(r, 2, value)
        ws.cell(r, 2).font = Font(name="Segoe UI", color="008000" if r in (4, 5, 6, 7) else C["ink"])
        ws.cell(r, 2).alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 28

    # ------------------------------------------------------------------ Epics
    epic_headers = [*("Epic", "Team", "Product", "Parent Link", "Parent Link Name", "Status",
                      "Epic Closed", "All Tickets Closed", "Closure Pending", "Total SP", "Done SP",
                      "Progress %", "Ticket Count", "Planned PI", "Discipline", "Sprint Movement",
                      "Summary", "Epic SP")]
    ew = [*(15, 23, 12, 16, 38, 18, 12, 16, 15, 12, 12, 12, 12, 13, 20, 24, 60, 12)]
    ews = setup_sheet("Epics", "Epics", "One row per Epic. Progress counts SP only from tickets whose Jira status is exactly Closed.", epic_headers, ew, C["purple"])
    ticket_rows = []
    for e in epics:
        tickets = list(e.get("tickets") or [])
        all_closed = bool(tickets) and all(_excel_exact_closed(t.get("status")) for t in tickets)
        epic_closed = _excel_exact_closed(e.get("status"))
        pending = all_closed and not epic_closed
        total = float(e.get("totalSp") or e.get("sp") or 0)
        closed_sp = sum(float(t.get("sp") or 0) for t in tickets if _excel_exact_closed(t.get("status")))
        progress = min(1.0, closed_sp / total) if total else (1.0 if epic_closed and all_closed else 0.0)
        progress = 0.99 if pending and progress >= 1 else progress
        codes = sorted({code for t in tickets for code in (t.get("sprints") or [])})
        ews.append([
            e.get("key", ""), e.get("team", ""), _portfolio_product(e),
            e.get("parentLink", ""), e.get("parentLinkName", ""), e.get("status", ""),
            epic_closed, all_closed, pending, total, closed_sp, None, len(tickets),
            e.get("ppi", ""), e.get("discipline", ""), ", ".join(codes), e.get("summary", ""),
            float(e.get("sp") or 0),
        ])
        er = ews.max_row
        ews.cell(er, 12, f"=IF(I{er},99%,IF(AND(G{er},H{er}),100%,IFERROR(K{er}/J{er},0)))")
        for t in tickets:
            ticket_rows.append([
                e.get("key", ""), t.get("key", ""), t.get("team") or "No team",
                _portfolio_product(e), e.get("parentLink", ""), t.get("status", ""),
                t.get("statusId", ""), t.get("assignee", ""),
                _excel_exact_closed(t.get("status")),
                float(t.get("sp") or 0), float(t.get("doneSp") or 0),
                t.get("discipline", ""), t.get("issueType", ""),
                ", ".join(t.get("sprints") or []), t.get("summary", ""),
            ])
    for row in range(5, ews.max_row + 1):
        ews.cell(row, 12).number_format = "0.0%"
        ews.cell(row, 10).number_format = ews.cell(row, 11).number_format = "0.0"
        if ews.cell(row, 9).value:
            ews.cell(row, 9).fill = PatternFill("solid", fgColor=C["soft_amber"])
            ews.cell(row, 9).font = Font(color=C["amber"], bold=True)
    if ews.max_row >= 5:
        tab = Table(displayName="EpicData", ref=f"A4:R{ews.max_row}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False)
        ews.add_table(tab)

    # ---------------------------------------------------------------- Tickets
    th = [*("Epic", "Ticket", "Team", "Product", "Parent Link", "Status", "Status ID", "Assignee",
            "Closed", "SP", "Done SP", "Discipline", "Issue Type", "Sprint(s)", "Summary")]
    tw = [*(15, 15, 23, 12, 16, 18, 12, 24, 10, 10, 10, 20, 15, 24, 60)]
    tws = setup_sheet("Tickets", "Child Tickets", "Only tickets attached to the in-scope Epics are stored. Tickets without an Epic are outside the PI portfolio.", th, tw, C["green"])
    for row in ticket_rows:
        tws.append(row)
    for row in range(5, tws.max_row + 1):
        tws.cell(row, 8).number_format = tws.cell(row, 9).number_format = "0.0"
    if tws.max_row >= 5:
        tab = Table(displayName="TicketData", ref=f"A4:O{tws.max_row}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False)
        tws.add_table(tab)

    # ------------------------------------------- Parent Link / Team aggregates
    parent_groups = {}
    team_groups = {}
    for e in epics:
        key = e.get("parentLink") or "No parent link"
        g = parent_groups.setdefault(key, {"name": e.get("parentLinkName") or key, "teams": set(),
                                           "epics": 0, "done": 0, "sp": 0.0, "doneSp": 0.0})
        g["teams"].add(e.get("team") or "No team")
        g["epics"] += 1
        g["done"] += 1 if (_excel_exact_closed(e.get("status")) and bool(e.get("tickets")) and all(_excel_exact_closed(t.get("status")) for t in (e.get("tickets") or []))) else 0
        total = float(e.get("totalSp") or e.get("sp") or 0)
        d = sum(float(t.get("sp") or 0) for t in (e.get("tickets") or []) if _excel_exact_closed(t.get("status")))
        g["sp"] += total
        g["doneSp"] += min(total, d)
        team = e.get("team") or "No team"
        q = team_groups.setdefault(team, {"parents": set(), "epics": 0, "done": 0, "sp": 0.0, "doneSp": 0.0})
        q["parents"].add(key)
        q["epics"] += 1
        q["done"] += 1 if (_excel_exact_closed(e.get("status")) and bool(e.get("tickets")) and all(_excel_exact_closed(t.get("status")) for t in (e.get("tickets") or []))) else 0
        q["sp"] += total
        q["doneSp"] += min(total, d)

    ph = [*("Parent Link", "Parent Link Name", "Product", "Teams", "Epics", "Done Epics", "Total SP", "Done SP", "Progress %")]
    pws = setup_sheet("Parent Links", "Parent Link Progress", "Complete Parent Link groups aggregated from all Epics and child tickets.", ph, [*(16, 42, 12, 45, 10, 12, 12, 12, 12)], C["orange"])
    for key, g in sorted(parent_groups.items()):
        product = "SPO4" if (str(key).upper().startswith("SPO4-") or str(g["name"]).upper().startswith("SPO4-")) else "Maintain"
        pws.append([key, g["name"], product, ", ".join(sorted(g["teams"])), g["epics"], g["done"], g["sp"], g["doneSp"], None])
        pr = pws.max_row
        pws.cell(pr, 9, f"=IFERROR(H{pr}/G{pr},0)")
    for r in range(5, pws.max_row + 1):
        pws.cell(r, 9).number_format = "0.0%"
        pws.cell(r, 7).number_format = pws.cell(r, 8).number_format = "0.0"
    if pws.max_row >= 5:
        tab = Table(displayName="ParentLinkData", ref=f"A4:I{pws.max_row}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True, showColumnStripes=False)
        pws.add_table(tab)

    teams_ws = setup_sheet("Teams", "Team Progress", "Portfolio totals grouped by Jira Team.", [*("Team", "Parent Links", "Epics", "Done Epics", "Total SP", "Done SP", "Progress %")], [*(24, 14, 10, 12, 12, 12, 12)], C["green"])
    for team, g in sorted(team_groups.items()):
        teams_ws.append([team, len(g["parents"]), g["epics"], g["done"], g["sp"], g["doneSp"], None])
        tr = teams_ws.max_row
        teams_ws.cell(tr, 7, f"=IFERROR(F{tr}/E{tr},0)")
    for r in range(5, teams_ws.max_row + 1):
        teams_ws.cell(r, 7).number_format = "0.0%"
        teams_ws.cell(r, 5).number_format = teams_ws.cell(r, 6).number_format = "0.0"
    if teams_ws.max_row >= 5:
        tab = Table(displayName="TeamData", ref=f"A4:G{teams_ws.max_row}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True, showColumnStripes=False)
        teams_ws.add_table(tab)

    # ------------------------------------------- hidden JSON payload sheets
    bws = wb.create_sheet("_BoardTickets")
    bws.sheet_state = "hidden"
    bws["A1"] = "JSON chunks"
    board_payload = json.dumps((data or {}).get("boardTickets") or [], ensure_ascii=False, separators=(",", ":"))
    board_chunk = 30000
    for i in range((len(board_payload) + board_chunk - 1) // board_chunk):
        bws.cell(i + 2, 1, board_payload[i * board_chunk:(i + 1) * board_chunk])

    # ------------------------------------------------------------ Commitment
    predictability = dict((data or {}).get("predictability") or {})
    predictability_rows = list(predictability.get("rows") or [])
    commitment_headers = [*("Team / board", "Board ID", "Sprint", "Sprint ID", "Final commitment", "Completed work", "Predictability", "Closed date")]
    commitment_widths = [*(28, 12, 16, 12, 20, 18, 16, 30)]
    cdata = setup_sheet("Commitment", "Sprint Predictability",
                        "Final commitment and Completed work per dashboard team and PI sprint. Predictability = Completed work / Final commitment.",
                        commitment_headers, commitment_widths, C["amber"])
    for item in predictability_rows:
        closed_value = item.get("closedDate") or ""
        if closed_value:
            try:
                closed_value = dt.datetime.fromisoformat(str(closed_value).replace("Z", "+00:00")).replace(tzinfo=None)
            except Exception:
                pass
        cdata.append([item.get("team") or "", item.get("boardId"), item.get("sprint") or "", item.get("sprintId"),
                      item.get("finalCommitment"), item.get("completedWork"), item.get("predictability"), closed_value])
        row = cdata.max_row
        cdata.cell(row, 5).number_format = "0.0#"
        cdata.cell(row, 6).number_format = "0.0#"
        cdata.cell(row, 7).number_format = "0.0%"
        cdata.cell(row, 8).number_format = "dd mmm yyyy, hh:mm"
    if cdata.max_row >= 5:
        tab = Table(displayName="CommitmentData", ref=f"A4:H{cdata.max_row}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True, showColumnStripes=False)
        cdata.add_table(tab)
    else:
        cdata.cell(5, 1, "No sprint predictability data is available. Check Force commitment refresh and click Sync.")
        cdata.merge_cells(start_row=5, start_column=1, end_row=5, end_column=8)
        cdata.cell(5, 1).font = Font(name="Segoe UI", size=10, color=C["amber"], italic=True)
        cdata.cell(5, 1).fill = PatternFill("solid", fgColor=C["soft_amber"])
        cdata.cell(5, 1).alignment = Alignment(vertical="center", wrap_text=True)
        cdata.row_dimensions[5].height = 32

    # Kept under the old sheet name so older workbooks stay readable.
    cws = wb.create_sheet("_Commitment")
    cws.sheet_state = "hidden"
    cws["A1"] = "JSON chunks"
    commitment_payload = json.dumps(predictability, ensure_ascii=False, separators=(",", ":"))
    for i in range((len(commitment_payload) + board_chunk - 1) // board_chunk):
        cws.cell(i + 2, 1, commitment_payload[i * board_chunk:(i + 1) * board_chunk])

    pws = wb.create_sheet("_Predictability")
    pws.sheet_state = "hidden"
    pws["A1"] = "JSON chunks"
    predictability_payload = json.dumps(predictability, ensure_ascii=False, separators=(",", ":"))
    for i in range((len(predictability_payload) + board_chunk - 1) // board_chunk):
        pws.cell(i + 2, 1, predictability_payload[i * board_chunk:(i + 1) * board_chunk])

    mws = wb.create_sheet("_Meta")
    mws.sheet_state = "hidden"
    mws.append(["Key", "Value"])
    for k, v in meta.items():
        mws.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v])

    # Turn every Jira key in the visible sheets into a clickable browse link.
    jira_base = str(meta.get("jiraUrl") or "").rstrip("/")
    jira_code_headers = {*frozenset({"Epic", "Ticket", "Parent Link"})}
    if jira_base:
        import re
        jira_key_pattern = re.compile("^[A-Z][A-Z0-9]+-\\d+$")
        for link_ws in wb.worksheets:
            if link_ws.sheet_state != "visible" or link_ws.max_row < 4:
                continue
            header_columns = [
                cell.column for cell in link_ws[4]
                if str(cell.value or "").strip() in jira_code_headers
            ]
            for col in header_columns:
                for row in range(5, link_ws.max_row + 1):
                    cell = link_ws.cell(row, col)
                    key = str(cell.value or "").strip()
                    if not jira_key_pattern.match(key):
                        continue
                    cell.hyperlink = f"{jira_base}/browse/{key}"
                    cell.style = "Hyperlink"

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass
    wb.save(path)
    return path


def read_portfolio_excel(path=None):
    """Rebuild dashboard data from the workbook generated by sync."""
    from openpyxl import load_workbook
    path = path or data_excel_path()
    wb = load_workbook(path, read_only=True, data_only=True)
    meta = {}
    if "_Meta" in wb.sheetnames:
        for key, value in list(wb["_Meta"].iter_rows(values_only=True))[1:]:
            if not key:
                continue
            if isinstance(value, str) and value[:1] in "[{":
                try:
                    value = json.loads(value)
                except Exception:
                    pass
            meta[str(key)] = value
    epics = []
    by = {}
    if "Epics" in wb.sheetnames:
        rows = list(wb["Epics"].iter_rows(min_row=4, values_only=True))
        headers = list(rows[0]) if rows else []
        for vals in rows[1:]:
            if not vals or not vals[0]:
                continue
            d = dict(zip(headers, vals))
            e = {
                "key": d.get("Epic") or "",
                "team": d.get("Team") or "No team",
                "parentLink": d.get("Parent Link") or "",
                "parentLinkName": d.get("Parent Link Name") or "",
                "status": d.get("Status") or "",
                "closed": bool(d.get("Epic Closed")),
                "sp": float(d.get("Epic SP") if d.get("Epic SP") not in (None, "") else (d.get("Total SP") or 0)),
                "totalSp": float(d.get("Total SP") or d.get("Epic SP") or 0),
                "doneSp": float(d.get("Done SP") or 0),
                "ppi": d.get("Planned PI") or "",
                "discipline": d.get("Discipline") or "",
                "summary": d.get("Summary") or "",
                "tickets": [],
            }
            epics.append(e)
            by[e["key"]] = e
    if "Tickets" in wb.sheetnames:
        rows = list(wb["Tickets"].iter_rows(min_row=4, values_only=True))
        headers = list(rows[0]) if rows else []
        for vals in rows[1:]:
            if not vals or not vals[0] or not vals[1]:
                continue
            d = dict(zip(headers, vals))
            e = by.get(str(d.get("Epic") or ""))
            if not e:
                continue
            sprints = [x.strip() for x in str(d.get("Sprint(s)") or "").split(",") if x.strip()]
            e["tickets"].append({
                "key": d.get("Ticket") or "",
                "epicKey": e["key"],
                "team": d.get("Team") or "No team",
                "status": d.get("Status") or "",
                "statusId": str(d.get("Status ID") or ""),
                "assignee": d.get("Assignee") or "",
                "closed": bool(d.get("Closed")),
                "sp": float(d.get("SP") or 0),
                "doneSp": float(d.get("Done SP") or 0),
                "discipline": d.get("Discipline") or "",
                "issueType": d.get("Issue Type") or "",
                "sprints": sprints,
                "summary": d.get("Summary") or "",
            })
    board_tickets = []
    commitment = {}
    predictability = {}
    if "_BoardTickets" in wb.sheetnames:
        parts = [str(row[0]) for row in wb["_BoardTickets"].iter_rows(min_row=2, max_col=1, values_only=True) if row[0] is not None]
        try:
            board_tickets = json.loads("".join(parts))
        except Exception:
            board_tickets = []
    if "_Commitment" in wb.sheetnames:
        parts = [str(row[0]) for row in wb["_Commitment"].iter_rows(min_row=2, max_col=1, values_only=True) if row[0] is not None]
        try:
            compatibility_payload = json.loads("".join(parts))
            if isinstance(compatibility_payload, dict) and "rows" in compatibility_payload:
                predictability = compatibility_payload
            else:
                commitment = compatibility_payload
        except Exception:
            commitment = {}
    if "_Predictability" in wb.sheetnames:
        parts = [str(row[0]) for row in wb["_Predictability"].iter_rows(min_row=2, max_col=1, values_only=True) if row[0] is not None]
        try:
            predictability = json.loads("".join(parts))
        except Exception:
            predictability = {}
    return {"meta": meta, "epics": epics, "boardTickets": board_tickets,
            "commitment": commitment, "predictability": predictability}


def _reuse_cached_ticket_progress(epics, board_tickets, cached_data):
    """Reuse sprint movement already stored in Excel during a fast Sync.

    Fresh Jira fields such as status, Team, SP and summary stay authoritative.
    Only sprint history is copied from the existing workbook. If Jira's fast
    response omits a previously known child or board ticket, the cached ticket
    is retained so the dashboard does not lose its working Kanban dataset.
    """
    cached_data = cached_data or {}
    cached_epics = list(cached_data.get("epics") or [])
    cached_board = list(cached_data.get("boardTickets") or [])

    cached_by_key = {}
    for epic in cached_epics:
        for ticket in (epic.get("tickets") or []):
            key = str(ticket.get("key") or "").strip()
            if not key:
                continue
            cached_by_key[key] = dict(ticket)
    for ticket in cached_board:
        key = str(ticket.get("key") or "").strip()
        if not key:
            continue
        current = cached_by_key.get(key, {})
        merged = dict(current)
        merged.update(ticket)
        if not (ticket.get("sprints") or []) and (current.get("sprints") or []):
            merged["sprints"] = list(current.get("sprints") or [])
        cached_by_key[key] = merged

    fresh_seen = set()
    epic_by_key = {str(e.get("key") or ""): e for e in (epics or [])}
    for epic in (epics or []):
        fresh_tickets = list(epic.get("tickets") or [])
        for ticket in fresh_tickets:
            key = str(ticket.get("key") or "").strip()
            if not key:
                continue
            fresh_seen.add(key)
            cached = cached_by_key.get(key) or {}
            if not cached.get("sprints"):
                continue
            ticket["sprints"] = list(cached.get("sprints") or [])

        # Keep child tickets Jira did not return in this fast response.
        existing = {str(t.get("key") or "") for t in fresh_tickets}
        for cached_epic in cached_epics:
            if str(cached_epic.get("key") or "") != str(epic.get("key") or ""):
                continue
            for cached_ticket in (cached_epic.get("tickets") or []):
                key = str(cached_ticket.get("key") or "").strip()
                if not key or key in existing:
                    continue
                fresh_tickets.append(dict(cached_ticket))
                existing.add(key)
                fresh_seen.add(key)
            break
        epic["tickets"] = fresh_tickets

    # Same treatment for the board/Kanban dataset.
    merged_board = []
    board_seen = set()
    for ticket in (board_tickets or []):
        key = str(ticket.get("key") or "").strip()
        if not key:
            continue
        cached = cached_by_key.get(key) or {}
        if cached.get("sprints"):
            ticket["sprints"] = list(cached.get("sprints") or [])
        merged_board.append(ticket)
        board_seen.add(key)
    for key, cached in cached_by_key.items():
        if key in board_seen or str(cached.get("epicKey") or "") not in epic_by_key:
            continue
        merged_board.append(dict(cached))
        board_seen.add(key)

    sprint_count = sum(1 for t in cached_by_key.values() if t.get("sprints"))
    print(f"[progress-cache] reused sprint movement for {sprint_count} cached ticket(s); "
          f"{sum(len(e.get('tickets') or []) for e in (epics or []))} child ticket(s) and "
          f"{len(merged_board)} board/Kanban ticket(s) available", flush=True)
    return epics, merged_board
